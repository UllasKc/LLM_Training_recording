"""
Generate the stakeholder brief (input 3) for each use case as a Word document.

Outputs:
    recording_workspace/UC1_Cross_Sell/Deliverables/input 3 - campaign brief.docx
    recording_workspace/UC2_Propensity_To_Pay/Deliverables/input 3 - collections strategy brief.docx

The brief carries the business context, constraints and acceptance criteria that the
expert must honour on camera. All names and numbers are fictional.

Usage:  python prep/generate_briefs.py
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
WS = ROOT / "recording_workspace"

UC1 = {
    "path": WS / "UC1_Cross_Sell" / "Deliverables" / "input 3 - campaign brief.docx",
    "title": "Cross-Sell Targeting Model - Brief for Q4 Campaign",
    "meta": [("From", "Head of CRM & Cross-Sell (Marketing)"), ("To", "Data Science team"),
             ("Date", "1 October 2026"), ("Decision needed by", "15 October 2026 Q4 campaign planning meeting")],
    "sections": [
        ("Background", [
            "Our campaigns offer an additional product (credit card, savings, personal loan or insurance) "
            "to existing customers. So far we have contacted customers at random and converted roughly "
            "1 in 9. We want to stop spending budget on customers who are very unlikely to buy.",
            "The Data Warehouse team has shared a snapshot of 20,000 customers as of 30 June 2026, together "
            "with whether each customer opened an additional product in the following 90 days (July - "
            "September). File: input 1 - customers.csv, definitions in input 2 - data dictionary.xlsx.",
            "Before we commit the Q4 budget we want evidence: if we had used a model on 30 June, how many "
            "more customers would we have converted than with random selection?",
        ]),
        ("What we need", [
            "1. A model that ranks customers by their likelihood to buy another product in the next 90 days, "
            "packaged so we can score the 30 September snapshot for Q4.",
            "2. A back-test target list: the 4,000 customers the model would have selected on 30 June, scored "
            "only with models that did not train on those customers, with expected and actual conversions "
            "and ROI compared to random selection.",
            "3. The main drivers of purchase in plain business language, so the creative team can tailor messaging.",
            "4. A short (3-slide) summary for the campaign planning meeting.",
        ]),
        ("Business constraints", [
            "Budget: 4,000 contacts per campaign. Cost per contact (incl. incentive): $4.",
            "Value: average 12-month margin per additional product sold: $150.",
            "Compliance: only customers with marketing_opt_in = 1 may be contacted.",
            "Compliance: gender must NOT be used as a model input (fair-treatment policy).",
            "Customer care rule: exclude customers with 2 or more complaints in the last 12 months from the target list.",
            "The model must only use information that is available on the snapshot date - it will be "
            "re-run on future snapshots before each campaign.",
        ]),
        ("Acceptance criteria", [
            "Performance on a hold-out set the model never saw: ROC-AUC of at least 0.78.",
            "Beats a simple logistic-regression baseline on PR-AUC (or the reason is explained).",
            "Top 20% of scored customers convert at least 2.5x the average rate (lift >= 2.5).",
            "No leakage: no field that is only populated after the purchase is used.",
            "Back-test list honours every compliance and customer-care rule above, has exactly 4,000 rows, "
            "and its actual conversion rate is at least 2x the eligible-customer average.",
            "Results can be reproduced by re-running one script end to end.",
        ]),
    ],
}

UC2 = {
    "path": WS / "UC2_Propensity_To_Pay" / "Deliverables" / "input 3 - collections strategy brief.docx",
    "title": "Collections Prioritisation - Propensity-to-Pay Brief",
    "meta": [("From", "Head of Collections Strategy"), ("To", "Data Science team"),
             ("Date", "18 August 2026"), ("Decision needed by", "September collections planning")],
    "sections": [
        ("Background", [
            "About 3,500 accounts (credit card, personal loan and BNPL) enter early collections every month. "
            "Our 25 agents can work about 1,300 accounts a month in depth (roughly 37% of intake). "
            "Today agents call accounts in order of highest days-past-due first.",
            "Cure rates have been falling through 2026, especially on BNPL, and we are missing our monthly "
            "collections target. Accounts that are likely to pay on their own should get an SMS/app nudge "
            "instead of an agent call, so agents can focus where they change the outcome.",
            "Attached: 20,000 accounts that entered collections between January and June 2026 with whether "
            "they paid the past-due amount in full within 30 days (input 1 - collections accounts.csv; "
            "definitions in input 2 - data dictionary.xlsx). Extracted 15 August 2026.",
        ]),
        ("What we need", [
            "1. A model that predicts each account's probability of paying in full within 30 days.",
            "2. A prioritised agent work list for the most recent cohort (June 2026) that fits agent "
            "capacity (37% of the cohort), ranked by expected unpaid amount = (1 - probability) x amount_due.",
            "3. The main drivers of payment, so the strategy team can adjust treatment paths.",
            "4. A short (3-slide) summary for September planning.",
        ]),
        ("Business constraints", [
            "Accounts are re-prioritised on day 7 after entering collections. Information captured in the "
            "first 7 days (contact attempts, right-party contact, promise to pay) may be used; anything "
            "captured later may not.",
            "The model will be used on future months, so it must be validated on later months than it was "
            "trained on (train on Jan-Apr, validate on May-Jun).",
            "Probabilities feed the expected-loss calculation, so they must be well calibrated, not just well ranked.",
            "Accounts with an open dispute or hardship flag always go to a specialist team - they stay in the "
            "model but are excluded from the agent work list.",
        ]),
        ("Acceptance criteria", [
            "Out-of-time (May-Jun) ROC-AUC of at least 0.78.",
            "Beats the current rule (highest days-past-due first) on expected unpaid amount captured within "
            "agent capacity, and beats a logistic-regression baseline on PR-AUC (or the reason is explained).",
            "Calibration checked on the out-of-time period (Brier score and reliability curve); "
            "recalibrated if predicted and actual pay rates differ materially.",
            "No leakage: no field captured after day 7 is used.",
            "Work list fits capacity, excludes dispute / hardship accounts, and is reproducible from one script.",
        ]),
    ],
}


def build(spec: dict) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(spec["title"], level=1)
    for label, value in spec["meta"]:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)
    for heading, paras in spec["sections"]:
        doc.add_heading(heading, level=2)
        for text in paras:
            numbered = text[0].isdigit()
            bullet = heading != "Background" and not numbered
            doc.add_paragraph(text, style="List Bullet" if bullet else None)
    spec["path"].parent.mkdir(parents=True, exist_ok=True)
    doc.save(spec["path"])
    print(f"wrote {spec['path']}")


if __name__ == "__main__":
    build(UC1)
    build(UC2)
