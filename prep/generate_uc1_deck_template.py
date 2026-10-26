"""
Generate input 4 - summary deck template.pptx for UC1: a 4-slide team template
(title + 3 content slides) whose orange [ ] placeholders the expert fills on camera
with the results and charts from notebook 03 / 04.

Needs python-pptx, which is NOT part of the recording environment:
    pip install python-pptx
    python prep/generate_uc1_deck_template.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "recording_workspace" / "UC1_Cross_Sell" / "Deliverables" / "1_Input" / "input 4 - summary deck template.pptx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
INK = RGBColor(0x26, 0x2B, 0x33)
MUTED = RGBColor(0x5F, 0x6B, 0x7A)
FILL = RGBColor(0xC2, 0x5B, 0x12)      # placeholder colour: anything orange still needs filling in
PANEL = RGBColor(0xF1, 0xF4, 0xF8)


def text_box(slide, x, y, w, h, runs, size=16, align=PP_ALIGN.LEFT, space_after=6):
    """runs: list of paragraphs; each paragraph is a list of (text, bold, colour)."""
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        for text, bold, colour in para:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = colour
            r.font.name = "Calibri"
    return tf


def header(slide, title, number):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    text_box(slide, Inches(0.6), Inches(0.35), Inches(11), Inches(0.8), [[(title, True, NAVY)]], size=30)
    text_box(slide, Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.35),
             [[("Cross-sell targeting model  ·  Data Science  ·  Internal", False, MUTED)]], size=10)
    text_box(slide, Inches(12.2), Inches(7.0), Inches(0.6), Inches(0.35), [[(str(number), False, MUTED)]],
             size=10, align=PP_ALIGN.RIGHT)


def kpi(slide, x, y, label, value_runs, note):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.8), Inches(1.55))
    box.fill.solid(); box.fill.fore_color.rgb = PANEL; box.line.fill.background()
    box.adjustments[0] = 0.08
    text_box(slide, x + Inches(0.2), y + Inches(0.12), Inches(3.4), Inches(0.4), [[(label, False, MUTED)]], size=13)
    text_box(slide, x + Inches(0.2), y + Inches(0.45), Inches(3.4), Inches(0.6), [value_runs], size=28)
    text_box(slide, x + Inches(0.2), y + Inches(1.08), Inches(3.4), Inches(0.4), [[(note, False, MUTED)]], size=12)


def bullets(slide, x, y, w, h, items, size=17):
    paras = []
    for item in items:
        runs = [("•  ", False, NAVY)]
        runs += item if isinstance(item, list) else [(item, False, INK)]
        paras.append(runs)
    text_box(slide, x, y, w, h, paras, size=size, space_after=10)


def image_placeholder(slide, x, y, w, h, label):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    box.fill.background()
    box.line.color.rgb = FILL; box.line.width = Pt(1.5); box.line.dash_style = 4  # dash
    tf = box.text_frame
    tf.text = label
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(14); r.font.color.rgb = FILL; r.font.bold = True


def F(text="[   ]"):
    return (text, True, FILL)


def T(text):
    return (text, False, INK)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 - title
    s = prs.slides.add_slide(blank)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background()
    white = RGBColor(0xFF, 0xFF, 0xFF)
    text_box(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.2), [[("Cross-Sell Targeting Model", True, white)]], size=44)
    text_box(s, Inches(0.9), Inches(3.5), Inches(11.5), Inches(0.8),
             [[("Back-test on the 30 June snapshot and Q4 recommendation", False, white)]], size=22)
    text_box(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(0.8),
             [[("Data Science  ·  For: Head of CRM & Cross-Sell  ·  Q4 planning, 15 October 2026", False,
                RGBColor(0xC9, 0xD6, 0xE8))]], size=14)

    # 2 - problem and approach
    s = prs.slides.add_slide(blank)
    header(s, "Problem and approach", 2)
    bullets(s, Inches(0.6), Inches(1.4), Inches(12), Inches(5.3), [
        "Random contacts convert ~1 in 9; budget is 4,000 contacts at $4 each, $150 margin per extra product",
        "Data: 20,000-customer snapshot (30 Jun 2026) with July–September purchase outcomes",
        [T("Data quality: removed "), F(), T(" duplicate rows (CRM merge); fixed mixed date formats, region labels and sentinel values")],
        "Removed 2 fields that are only populated after a purchase (leakage) and gender (fair-treatment policy)",
        "Logistic-regression baseline → gradient boosting; evaluated on an untouched 20% hold-out",
        "Back-test: every customer scored by a model that never saw them (5-fold out-of-fold)",
    ])

    # 3 - model results
    s = prs.slides.add_slide(blank)
    header(s, "Model results (hold-out set)", 3)
    kpi(s, Inches(0.6), Inches(1.4), "ROC-AUC", [F()], "target ≥ 0.78")
    kpi(s, Inches(0.6), Inches(3.15), "PR-AUC", [F(), T("  vs "), F(), T(" baseline")], "random = 0.116")
    kpi(s, Inches(0.6), Inches(4.9), "Top-20% lift", [F("[   ]×")], "target ≥ 2.5×")
    image_placeholder(s, Inches(4.8), Inches(1.4), Inches(7.9), Inches(5.05),
                      "Insert chart: 3_Intermediate / artifact 7 - evaluation charts / lift_by_decile.png")

    # 4 - back-test and recommendation
    s = prs.slides.add_slide(blank)
    header(s, "Back-test and Q4 recommendation", 4)
    kpi(s, Inches(0.6), Inches(1.4), "Conversions (4,000 contacted)", [F(), T("  vs "), F(), T(" random")], "actual Jul–Sep outcomes")
    kpi(s, Inches(0.6), Inches(3.15), "Profit", [F("$[   ]k"), T(" vs "), F("$[   ]k")], "conversions × $150 − $16k")
    kpi(s, Inches(0.6), Inches(4.9), "Conversion rate", [F("[   ]%"), T(" vs "), F("[   ]%")], "target list vs eligible average")
    text_box(s, Inches(4.8), Inches(1.35), Inches(7.9), Inches(0.5), [[("What drives purchase", True, NAVY)]], size=18)
    bullets(s, Inches(4.8), Inches(1.85), Inches(7.9), Inches(2.4), [
        [F("[Driver 1 from SHAP]")], [F("[Driver 2 from SHAP]")], [F("[Driver 3 from SHAP]")],
    ], size=16)
    text_box(s, Inches(4.8), Inches(4.2), Inches(7.9), Inches(0.5), [[("Recommendation", True, NAVY)]], size=18)
    bullets(s, Inches(4.8), Inches(4.7), Inches(7.9), Inches(2.2), [
        "Select the Q4 list with the packaged model on the 30-Sep snapshot (opted-in, < 2 complaints)",
        "Tailor messaging to the drivers above; keep a 10% random control group to measure uplift",
        "Re-validate each quarter: compare expected vs actual conversions",
    ], size=15)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
